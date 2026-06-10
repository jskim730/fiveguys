"""Build the COSE471 presentation deck (python-pptx). Run: python report/build_slides.py"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT = HERE / "COSE471_slides.pptx"

# ── palette ──
DARK = RGBColor(0x14, 0x21, 0x3D)
LIGHT = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
RED = RGBColor(0xC0, 0x39, 0x2B)     # STYLE7
GREEN = RGBColor(0x1B, 0x78, 0x37)   # STYLO
GOLD = RGBColor(0xE6, 0xA8, 0x17)
CARD = RGBColor(0xEC, 0xEF, 0xF3)
HEAD, BODY = "Trebuchet MS", "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = 13.333, 7.5


def slide(bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        if "after" in para:
            p.space_after = Pt(para["after"])
        if "before" in para:
            p.space_before = Pt(para["before"])
        if "line" in para:
            p.line_spacing = para["line"]
        for run in para["runs"]:
            r = p.add_run()
            r.text = run["t"]
            f = r.font
            f.size = Pt(run.get("size", 16))
            f.bold = run.get("bold", False)
            f.italic = run.get("italic", False)
            f.name = run.get("font", BODY)
            f.color.rgb = run.get("color", INK)
    return tb


def box(s, x, y, w, h, fill, line=None, radius=True, shadow=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def pic(s, name, bx, by, bw, bh):
    """Fit image inside a box, preserving aspect ratio, centered."""
    p = FIG / name
    iw, ih = Image.open(p).size
    ar = iw / ih
    w = bw
    h = w / ar
    if h > bh:
        h = bh
        w = h * ar
    x = bx + (bw - w) / 2
    y = by + (bh - h) / 2
    s.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(w), Inches(h))


def footer(s, n, label, dark=False):
    c = ICE if dark else MUTED
    text(s, 0.6, H - 0.45, 12.1, 0.3,
         [{"runs": [{"t": f"COSE471 Term Project   ·   {label}", "size": 10, "color": c, "font": BODY}]}])
    text(s, 12.0, H - 0.45, 0.8, 0.3,
         [{"runs": [{"t": f"{n}/11", "size": 10, "color": c, "font": BODY}]}], align=PP_ALIGN.RIGHT)


def title(s, t, sub=None):
    text(s, 0.6, 0.42, 12.1, 0.8,
         [{"runs": [{"t": t, "size": 30, "bold": True, "color": INK, "font": HEAD}]}])
    if sub:
        text(s, 0.62, 1.16, 12.1, 0.4,
             [{"runs": [{"t": sub, "size": 15, "color": MUTED, "font": BODY}]}])


def takeaway(s, x, y, w, h, accent, lines):
    box(s, x, y, w, h, CARD)
    box(s, x, y, 0.09, h, accent, radius=False)
    text(s, x + 0.28, y + 0.14, w - 0.5, h - 0.28, lines, anchor=MSO_ANCHOR.MIDDLE)


# ════════════════════════════════════════════════════════════════
# 1 — TITLE (dark)
s = slide(DARK)
box(s, 0, 0, 0.22, H, RED, radius=False)
box(s, 0.22, 0, 0.18, H, GREEN, radius=False)
text(s, 1.0, 1.5, 11.3, 0.4,
     [{"runs": [{"t": "COSE471  ·  DATA SCIENCE  ·  KOREA UNIVERSITY", "size": 15, "bold": True, "color": GOLD, "font": BODY}]}])
text(s, 0.95, 2.2, 11.6, 2.2, [
    {"runs": [{"t": "Surface Style is Coachable.", "size": 50, "bold": True, "color": WHITE, "font": HEAD}], "after": 4},
    {"runs": [{"t": "The Signature is Not.", "size": 50, "bold": True, "color": ICE, "font": HEAD}]},
])
text(s, 1.0, 4.55, 11.0, 0.7,
     [{"runs": [{"t": "Can an LLM imitate a student's writing style once the content is held fixed?",
                 "size": 19, "italic": True, "color": ICE, "font": BODY}]}])
text(s, 1.0, 6.4, 11.0, 0.5,
     [{"runs": [{"t": "Team [nickname]   ·   Members A · B · C · D", "size": 14, "color": MUTED, "font": BODY}]}])

# ════════════════════════════════════════════════════════════════
# 2 — PROBLEM (light)
s = slide(LIGHT)
title(s, "“Looks human” ≠ “is human”")
text(s, 0.62, 1.7, 6.5, 4.5, [
    {"runs": [{"t": "AI-text detectors usually confuse two different things:", "size": 18, "color": INK, "font": BODY}], "after": 14},
    {"runs": [{"t": "WHAT ", "size": 18, "bold": True, "color": GREEN, "font": BODY},
              {"t": "is said (content)  vs.  ", "size": 18, "color": INK, "font": BODY},
              {"t": "HOW ", "size": 18, "bold": True, "color": RED, "font": BODY},
              {"t": "it is said (style).", "size": 18, "color": INK, "font": BODY}], "after": 14},
    {"runs": [{"t": "We hold ", "size": 18, "color": INK, "font": BODY},
              {"t": "content fixed", "size": 18, "bold": True, "color": GREEN, "font": BODY},
              {"t": " and coach ", "size": 18, "color": INK, "font": BODY},
              {"t": "only the style", "size": 18, "bold": True, "color": RED, "font": BODY},
              {"t": " — then ask which part of “human” an LLM cannot fake.", "size": 18, "color": INK, "font": BODY}]},
])
box(s, 7.6, 1.85, 5.1, 1.85, CARD)
box(s, 7.6, 1.85, 0.12, 1.85, GREEN, radius=False)
text(s, 7.95, 2.05, 4.6, 1.5, [
    {"runs": [{"t": "CONTENT", "size": 16, "bold": True, "color": GREEN, "font": HEAD}], "after": 4},
    {"runs": [{"t": "Held FIXED — each LLM answer must cover that student's keywords.", "size": 15, "color": INK, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
box(s, 7.6, 4.0, 5.1, 1.85, CARD)
box(s, 7.6, 4.0, 0.12, 1.85, RED, radius=False)
text(s, 7.95, 4.2, 4.6, 1.5, [
    {"runs": [{"t": "STYLE", "size": 16, "bold": True, "color": RED, "font": HEAD}], "after": 4},
    {"runs": [{"t": "The ONLY variable — coached harder each round (R0 → R2).", "size": 15, "color": INK, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2, "Motivation")

# ════════════════════════════════════════════════════════════════
# 3 — DATA (light)
s = slide(LIGHT)
title(s, "Real in-class discussion data", "Written by our own class — mixed languages, uneven length, genuine messiness.")
stats = [("5", "topics"), ("9", "questions"), ("542", "responses"), ("69", "students"), ("421", "English-original\n(analysis set)")]
sx = 0.62
for num, lab in stats:
    box(s, sx, 1.75, 2.34, 1.25, CARD)
    text(s, sx, 1.9, 2.34, 0.7, [{"runs": [{"t": num, "size": 32, "bold": True, "color": DARK, "font": HEAD}]}], align=PP_ALIGN.CENTER)
    text(s, sx + 0.1, 2.55, 2.14, 0.45, [{"runs": [{"t": lab, "size": 12, "color": MUTED, "font": BODY}]}], align=PP_ALIGN.CENTER)
    sx += 2.44
pic(s, "fig6_data_overview.png", 0.62, 3.2, 12.1, 3.7)
footer(s, 3, "Data understanding")

# ════════════════════════════════════════════════════════════════
# 4 — METHOD (light)
s = slide(LIGHT)
title(s, "Fix the content, coach only the style")
# pipeline chevrons
text(s, 0.62, 1.7, 3.0, 0.9, [
    {"runs": [{"t": "keywords", "size": 16, "bold": True, "color": GREEN, "font": HEAD}], "after": 2},
    {"runs": [{"t": "content fixed,\nall rounds", "size": 12, "color": MUTED, "font": BODY}]}], anchor=MSO_ANCHOR.MIDDLE)
rounds = [("R0", "naive", RGBColor(0xF2, 0xC4, 0xBE)), ("R1", "light", RGBColor(0xE0, 0x84, 0x77)), ("R2", "aggressive", RED)]
rx = 3.0
for code, lab, col in rounds:
    ch = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(rx), Inches(1.75), Inches(2.5), Inches(0.95))
    ch.fill.solid(); ch.fill.fore_color.rgb = col; ch.line.fill.background(); ch.shadow.inherit = False
    tcol = WHITE if code == "R2" else INK
    tf = ch.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f"{code}  "; r.font.bold = True; r.font.size = Pt(18); r.font.name = HEAD; r.font.color.rgb = tcol
    r2 = p.add_run(); r2.text = lab; r2.font.size = Pt(13); r2.font.name = BODY; r2.font.color.rgb = tcol
    rx += 2.25
text(s, 3.0, 2.78, 7.5, 0.4, [{"runs": [{"t": "style coaching only — persona, temperature fixed", "size": 12, "italic": True, "color": MUTED, "font": BODY}]}], align=PP_ALIGN.CENTER)

# judges
box(s, 0.62, 3.7, 6.0, 1.5, CARD)
box(s, 0.62, 3.7, 0.12, 1.5, RED, radius=False)
text(s, 0.95, 3.85, 5.5, 1.2, [
    {"runs": [{"t": "STYLE7", "size": 16, "bold": True, "color": RED, "font": HEAD}, {"t": "  —  targeted detector", "size": 13, "color": MUTED, "font": BODY}], "after": 3},
    {"runs": [{"t": "7 surface features the prompt explicitly coaches (length, 1st-person, hedging…).", "size": 13, "color": INK, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
box(s, 6.9, 3.7, 5.8, 1.5, CARD)
box(s, 6.9, 3.7, 0.12, 1.5, GREEN, radius=False)
text(s, 7.23, 3.85, 5.3, 1.2, [
    {"runs": [{"t": "STYLO", "size": 16, "bold": True, "color": GREEN, "font": HEAD}, {"t": "  —  independent judge", "size": 13, "color": MUTED, "font": BODY}], "after": 3},
    {"runs": [{"t": "function words + char n-grams + punctuation. Never targeted. Disjoint from STYLE7.", "size": 13, "color": INK, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
takeaway(s, 0.62, 5.45, 12.1, 1.15, GOLD, [
    {"runs": [{"t": "Two course techniques:  ", "size": 15, "bold": True, "color": INK, "font": BODY},
              {"t": "① Classification", "size": 15, "bold": True, "color": DARK, "font": BODY},
              {"t": " (per-question, leakage-proof AUC)   +   ", "size": 15, "color": INK, "font": BODY},
              {"t": "② Outlier Detection", "size": 15, "bold": True, "color": DARK, "font": BODY},
              {"t": " (LOF).   Generator GPT-5-mini ≠ judge.", "size": 15, "color": INK, "font": BODY}]},
])
footer(s, 4, "Methodology")

# ════════════════════════════════════════════════════════════════
# 5 — PER-QUESTION (light)
s = slide(LIGHT)
title(s, "Style encodes topic — so we never pool")
pic(s, "fig5_topic_confound.png", 0.7, 1.7, 6.2, 5.0)
takeaway(s, 7.4, 2.6, 5.3, 2.6, RED, [
    {"runs": [{"t": "Using humans only, style predicts WHICH question a response answers:", "size": 16, "color": INK, "font": BODY}], "after": 10},
    {"runs": [{"t": "STYLE7 AUC 0.82   ·   STYLO AUC 0.96", "size": 17, "bold": True, "color": DARK, "font": BODY}], "after": 10},
    {"runs": [{"t": "Features carry topic, not just authorship → we evaluate strictly per question (no topic leakage).", "size": 15, "color": INK, "font": BODY}]},
])
footer(s, 5, "Methodology check")

# ════════════════════════════════════════════════════════════════
# 6 — TELL + OVERSHOOT (light)
s = slide(LIGHT)
title(s, "The LLM “tell” — and the overshoot")
pic(s, "fig2_style7_feature_trajectory.png", 0.55, 1.55, 8.1, 5.3)
takeaway(s, 9.1, 1.7, 3.7, 1.5, RED, [
    {"runs": [{"t": "Naive R0", "size": 14, "bold": True, "color": RED, "font": HEAD}], "after": 3},
    {"runs": [{"t": "~6× longer, polished, ~zero hedging, impersonal.", "size": 13.5, "color": INK, "font": BODY}]}])
takeaway(s, 9.1, 3.35, 3.7, 1.5, GREEN, [
    {"runs": [{"t": "Coached R2", "size": 14, "bold": True, "color": GREEN, "font": HEAD}], "after": 3},
    {"runs": [{"t": "Overshoots humans on every feature (crosses the dashed line).", "size": 13.5, "color": INK, "font": BODY}]}])
takeaway(s, 9.1, 5.0, 3.7, 1.65, GOLD, [
    {"runs": [{"t": "Surprise (data-driven)", "size": 14, "bold": True, "color": DARK, "font": HEAD}], "after": 3},
    {"runs": [{"t": "Humans hedge MORE than naive GPT — our plan's assumption was backwards.", "size": 13.5, "color": INK, "font": BODY}]}])
footer(s, 6, "Results")

# ════════════════════════════════════════════════════════════════
# 7 — HEADLINE (light, emphasized)
s = slide(LIGHT)
title(s, "Surface breaks, signature holds", "Frozen R0 classifier scored on held-out humans, per question.")
pic(s, "fig1_auc_trajectory.png", 0.6, 1.65, 7.4, 5.2)
box(s, 8.3, 2.0, 4.4, 1.95, CARD)
box(s, 8.3, 2.0, 0.12, 1.95, RED, radius=False)
text(s, 8.6, 2.1, 4.0, 1.75, [
    {"runs": [{"t": "STYLE7 @ R2", "size": 15, "bold": True, "color": RED, "font": HEAD}], "after": 2},
    {"runs": [{"t": "0.08", "size": 40, "bold": True, "color": RED, "font": HEAD}], "after": 0},
    {"runs": [{"t": "surface coached away — inverted past humans", "size": 12.5, "color": MUTED, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
box(s, 8.3, 4.15, 4.4, 1.95, CARD)
box(s, 8.3, 4.15, 0.12, 1.95, GREEN, radius=False)
text(s, 8.6, 4.25, 4.0, 1.75, [
    {"runs": [{"t": "STYLO @ R2", "size": 15, "bold": True, "color": GREEN, "font": HEAD}], "after": 2},
    {"runs": [{"t": "0.78", "size": 40, "bold": True, "color": GREEN, "font": HEAD}], "after": 0},
    {"runs": [{"t": "independent judge still detects the LLM", "size": 12.5, "color": MUTED, "font": BODY}]},
], anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.3, 6.25, 4.4, 0.5, [{"runs": [{"t": "Same answers — opposite verdicts.", "size": 15, "italic": True, "bold": True, "color": DARK, "font": BODY}]}], align=PP_ALIGN.CENTER)
footer(s, 7, "Results — headline")

# ════════════════════════════════════════════════════════════════
# 8 — TWO TECHNIQUES (light)
s = slide(LIGHT)
title(s, "Two methods, one verdict")
pic(s, "fig3_two_technique_convergence.png", 0.6, 1.65, 6.0, 4.0)
pic(s, "fig4_stylo_pca_Q5.png", 6.9, 1.65, 5.9, 4.0)
takeaway(s, 0.62, 5.95, 12.1, 0.9, GREEN, [
    {"runs": [{"t": "Classifier (STYLO AUC 0.78) and Outlier-LOF (59% of humans still outside the LLM cloud) — distance-based, not a re-run — agree the signature survives imitation.",
               "size": 15, "color": INK, "font": BODY}]},
])
footer(s, 8, "Results — two techniques")

# ════════════════════════════════════════════════════════════════
# 9 — ROBUSTNESS (light)
s = slide(LIGHT)
title(s, "Not an artifact")
pic(s, "fig8_robustness.png", 0.7, 1.7, 6.4, 5.0)
takeaway(s, 7.6, 2.3, 5.1, 3.5, GOLD, [
    {"runs": [{"t": "The pattern holds across 4 data definitions:", "size": 16, "color": INK, "font": BODY}], "after": 10},
    {"runs": [{"t": "English-only (0.78) ≈ translation-included (0.76)", "size": 15, "bold": True, "color": DARK, "font": BODY}], "after": 4},
    {"runs": [{"t": "→ not a machine-translation artifact.", "size": 14, "color": MUTED, "font": BODY}], "after": 12},
    {"runs": [{"t": "At R2 human (41w) ≈ LLM (38w)", "size": 15, "bold": True, "color": DARK, "font": BODY}], "after": 4},
    {"runs": [{"t": "→ the residual 0.78 is not a length artifact.", "size": 14, "color": MUTED, "font": BODY}]},
])
footer(s, 9, "Results — robustness")

# ════════════════════════════════════════════════════════════════
# 10 — CONCLUSION (dark)
s = slide(DARK)
box(s, 0, 0, 0.22, H, RED, radius=False)
box(s, 0.22, 0, 0.18, H, GREEN, radius=False)
text(s, 1.0, 0.9, 11.6, 0.6, [{"runs": [{"t": "The signature is the human part", "size": 34, "bold": True, "color": WHITE, "font": HEAD}]}])
text(s, 1.0, 1.95, 11.4, 1.6, [
    {"runs": [{"t": "When content is fixed, prompt-level coaching erases the LLM's ", "size": 21, "color": ICE, "font": BODY},
              {"t": "surface tells", "size": 21, "bold": True, "color": RED, "font": BODY},
              {"t": " — but cannot erase its ", "size": 21, "color": ICE, "font": BODY},
              {"t": "stylometric signature", "size": 21, "bold": True, "color": GREEN, "font": BODY},
              {"t": ".", "size": 21, "color": ICE, "font": BODY}]},
])
nums = [("0.08", "STYLE7 @ R2", RED), ("0.78", "STYLO @ R2", GREEN), ("59%", "humans outside\nLLM cloud", GOLD)]
nx = 1.0
for num, lab, col in nums:
    box(s, nx, 3.85, 3.6, 1.85, RGBColor(0x1C, 0x2C, 0x4D))
    text(s, nx, 4.05, 3.6, 0.9, [{"runs": [{"t": num, "size": 44, "bold": True, "color": col, "font": HEAD}]}], align=PP_ALIGN.CENTER)
    text(s, nx + 0.2, 5.05, 3.2, 0.55, [{"runs": [{"t": lab, "size": 13, "color": ICE, "font": BODY}]}], align=PP_ALIGN.CENTER)
    nx += 3.85
text(s, 1.0, 6.45, 11.4, 0.7, [{"runs": [{"t": "Limitations: STYLO also drops (partial, not immune)  ·  R2 is a caricature (overshoot)  ·  ESL-vs-native register → future work.",
                                          "size": 13, "italic": True, "color": MUTED, "font": BODY}]}])

# ════════════════════════════════════════════════════════════════
# 11 — THANKS (dark)
s = slide(DARK)
text(s, 0, 2.4, W, 1.0, [{"runs": [{"t": "Thank you", "size": 52, "bold": True, "color": WHITE, "font": HEAD}]}], align=PP_ALIGN.CENTER)
text(s, 0, 3.7, W, 0.6, [{"runs": [{"t": "Questions & Discussion", "size": 22, "italic": True, "color": ICE, "font": BODY}]}], align=PP_ALIGN.CENTER)
text(s, 0, 5.4, W, 0.5, [{"runs": [{"t": "Reproducible: config.yaml  →  generate → classify → round_eval → outlier_lof → figures",
                                     "size": 13, "color": MUTED, "font": BODY}]}], align=PP_ALIGN.CENTER)
text(s, 0, 6.0, W, 0.5, [{"runs": [{"t": "Team [nickname]  ·  Members A · B · C · D", "size": 13, "color": MUTED, "font": BODY}]}], align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print("saved", OUT)
